import os
from docx import Document
from docx.shared import Pt
import re

def convert_md_to_docx(md_path, docx_path):
    document = Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        # Headers
        if line.startswith('# '):
            document.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            document.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            document.add_heading(line[4:], level=3)
        elif line.startswith('---'):
            pass # Ignore horizontal rules
        # Tables (Very basic handling - just print as text for now)
        elif line.startswith('|'):
            p = document.add_paragraph(line)
        # Lists
        elif line.startswith('- ') or line.startswith('* '):
            # Clean up bold syntax
            text = line[2:].replace('**', '')
            document.add_paragraph(text, style='List Bullet')
        elif re.match(r'^\d+\. ', line):
            text = re.sub(r'^\d+\. ', '', line).replace('**', '')
            document.add_paragraph(text, style='List Number')
        else:
            # Clean up bold syntax
            text = line.replace('**', '')
            document.add_paragraph(text)
            
    document.save(docx_path)
    print(f"Successfully converted to {docx_path}")

import sys

if __name__ == "__main__":
    if len(sys.argv) == 3:
        md_file = sys.argv[1]
        docx_file = sys.argv[2]
    else:
        md_file = os.path.join("documents", "review_1_report.md")
        docx_file = os.path.join("documents", "review_1_report.docx")
    convert_md_to_docx(md_file, docx_file)

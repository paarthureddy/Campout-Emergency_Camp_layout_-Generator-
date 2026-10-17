"""
Generate formatted Microsoft Word document (.docx) for Lab Evaluation 2 Report.
Includes styled headings, summary tables, and embedded high-resolution figures.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def create_report_docx():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("LAB EVALUATION 2: GAN MODEL WITH OTHER DEEP LEARNING MODEL")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(26, 54, 93) # Deep Blue

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Comparative Benchmarking & Product Testing in the ReliefPlan AI Case Study\nCourse: 23CSE473 – Neural Networks and Deep Learning (10 Marks)")
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(74, 85, 104)

    # Metadata Table
    table_meta = doc.add_table(rows=5, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Course & Evaluation", "23CSE473 – Neural Networks and Deep Learning | Lab Evaluation 2 (8-Oct-2026)"),
        ("Faculty In-Charge", "Dr. S. Jenita Bagyammal"),
        ("Case Study & Product", "ReliefPlan AI: AI-Powered Disaster Relief Camp Planning"),
        ("Paarthu Reddy (Team Lead)", "Roll No: CB.SC.U4CSE23639 | Model: Pix2Pix Conditional GAN (cGAN)"),
        ("Shivam Kharat (Researcher)", "Roll No: CB.SC.U4CSE23445 | Model: Cycle-Consistent GAN (CycleGAN)")
    ]

    for row_idx, (col1, col2) in enumerate(meta_info):
        row = table_meta.rows[row_idx]
        cell_a, cell_b = row.cells[0], row.cells[1]
        cell_a.text = col1
        cell_b.text = col2
        cell_a.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell_a, "EBF8FF")
        set_cell_background(cell_b, "F7FAFC")
        cell_a.width = Inches(2.2)
        cell_b.width = Inches(4.5)

    doc.add_paragraph() # Spacing

    # 1. Executive Summary
    h1 = doc.add_heading("1. Executive Summary & Problem Context", level=1)
    h1.style.font.color.rgb = RGBColor(26, 54, 93)

    p = doc.add_paragraph(
        "Natural disasters such as earthquakes, tsunamis, and extreme floods displace hundreds of thousands of people "
        "within hours. Rapid humanitarian shelter deployment is essential to minimize casualties and prevent disease outbreaks. "
        "Currently, disaster response teams manually design camp layouts using static paper handbooks and GIS software—a manual process "
        "taking days that frequently leads to hazardous overcrowding, inadequate emergency road access, and severe non-compliance with "
        "UNHCR (United Nations High Commissioner for Refugees) safety standards."
    )
    p = doc.add_paragraph(
        "ReliefPlan AI automates this pipeline end-to-end. Our ongoing project implementation utilizes supervised deep learning "
        "(DeepLabV3+ with Atrous Spatial Pyramid Pooling and baseline U-Net) to segment satellite imagery into 7 terrain classes, "
        "followed by a multi-objective constraint-based spatial solver. For Lab Evaluation 2, we advance from discriminative classification "
        "to Generative Adversarial Networks (GANs), investigating direct satellite-to-blueprint generation across both paired (Pix2Pix) "
        "and unpaired (CycleGAN) paradigms for students Paarthu Reddy and Shivam Kharat."
    )

    # 2. Model Formulations
    h1 = doc.add_heading("2. Theoretical Formulation & Architecture of Models", level=1)
    h1.style.font.color.rgb = RGBColor(26, 54, 93)

    # 2.1 Paarthu Reddy
    h2 = doc.add_heading("2.1 Paarthu Reddy (CB.SC.U4CSE23639) — Pix2Pix Conditional GAN (cGAN)", level=2)
    h2.style.font.color.rgb = RGBColor(43, 108, 176)
    doc.add_paragraph(
        "• Motivation & Case Study Function: When historical disaster records provide paired satellite-to-blueprint layout instances, "
        "Pix2Pix provides conditional supervised adversarial generation. The Generator is conditioned on input satellite image x.\n"
        "• Objective Function: L_Pix2Pix = argmin_G max_D L_cGAN(G, D) + λ L_L1(G), where λ = 100.\n"
        "  - L_cGAN(G, D) = E[log D(x, y)] + E[log(1 - D(x, G(x)))]\n"
        "  - L_L1(G) = E[||y - G(x)||_1]\n"
        "• Generator Architecture: U-Net 256 featuring 8 downsampling and 8 upsampling stages with skip connections and Instance Normalization. "
        "Total Generator Parameters: 54,413,955 (~54.41M).\n"
        "• Discriminator Architecture: 70x70 PatchGAN Markovian discriminator conditioned on concatenated 6-channel [x, y] inputs. "
        "Total Discriminator Parameters: 2,769,089 (~2.77M).\n"
        "• Model Total Parameters: 57,183,044 (~57.18M) | Disk Footprint: 218.14 MB."
    )

    # 2.2 Shivam Kharat
    h2 = doc.add_heading("2.2 Shivam Kharat (CB.SC.U4CSE23445) — Cycle-Consistent GAN (CycleGAN)", level=2)
    h2.style.font.color.rgb = RGBColor(43, 108, 176)
    doc.add_paragraph(
        "• Motivation & Case Study Function: In emergency operations, newly struck disaster areas almost never have paired 1:1 ground truth "
        "camp layout blueprints. CycleGAN resolves this fundamental data bottleneck by learning unpaired domain translation between Domain X "
        "(Satellite Terrains) and Domain Y (UNHCR Camp Blueprints).\n"
        "• Objective Function: L_CycleGAN = L_GAN(G_AB, D_B) + L_GAN(G_BA, D_A) + λ_cyc * L_cyc + λ_id * L_id, with λ_cyc=10, λ_id=5.\n"
        "  - Forward & Backward Cycle Consistency: L_cyc = E[||G_BA(G_AB(x)) - x||_1] + E[||G_AB(G_BA(y)) - y||_1]\n"
        "  - Identity Mapping Loss: L_id = E[||G_AB(y) - y||_1] + E[||G_BA(x) - x||_1]\n"
        "• Generator Architecture: Dual ResNet-6 generators with 6 residual bottleneck blocks, reflection padding, and Instance Normalization. "
        "Parameters per Generator: 7,841,411 (~7.84M).\n"
        "• Discriminator Architecture: Dual 70x70 PatchGAN discriminators (5,532,034 parameters combined).\n"
        "• Model Total Parameters: 21,214,856 (~21.21M) | Disk Footprint: 80.93 MB."
    )

    # 2.3 Ongoing Project Models
    h2 = doc.add_heading("2.3 Ongoing Deep Learning Models in ReliefPlan AI", level=2)
    h2.style.font.color.rgb = RGBColor(43, 108, 176)
    doc.add_paragraph(
        "• U-Net (Baseline Ongoing): 4-level encoder-decoder with skip connections, 31,038,023 parameters (~31.04M), 67.6 ms GPU latency. "
        "Trains with Cross-Entropy + Dice Loss for 7-class terrain segmentation.\n"
        "• DeepLabV3+ with ASPP (Proposed Ongoing): Lightweight ResNet backbone with Atrous Spatial Pyramid Pooling (dilation rates 6, 12, 18). "
        "Requires only 3,935,783 parameters (~3.94M, 87.3% smaller than U-Net) with an ultra-fast GPU inference latency of 5.79 ms (172.7 FPS)."
    )

    # 3. Testing Aspects Done with the Product
    h1 = doc.add_heading("3. Testing Aspects Done with the Product in the Case Study", level=1)
    h1.style.font.color.rgb = RGBColor(26, 54, 93)
    doc.add_paragraph(
        "To evaluate model viability within the ReliefPlan AI system, a comprehensive five-pillar testing battery was executed on the LoveDA dataset:\n"
        "1. Generative Quality & Structural Fidelity Testing: Structural Similarity Index (SSIM), Peak Signal-to-Noise Ratio (PSNR in dB), "
        "and Laplacian Edge Variance measuring boundary sharpness of roads and shelter perimeters.\n"
        "2. Computational & Profiling Benchmarking: Model parameters, disk footprint, GPU inference latency (RTX 4060 GPU), CPU latency, and throughput (FPS).\n"
        "3. UNHCR Humanitarian Planning Compliance Testing: Automated testing of the 2-meter shelter firebreak rule, 200-meter water access radius, "
        "50-meter latrine proximity, and hazard zone safety avoidance rate (% of shelters placed strictly away from lakes and flood lines).\n"
        "4. Adversarial Noise Perturbation Testing: Injected additive Gaussian noise (σ = 0.15) to evaluate degradation under sensor noise.\n"
        "5. Adversarial Dynamics Testing: Monitored Generator-Discriminator loss balance to verify convergence and prevent mode collapse."
    )

    # 4. Master Comparative Results
    h1 = doc.add_heading("4. Master Comparative Experimental Results", level=1)
    h1.style.font.color.rgb = RGBColor(26, 54, 93)

    # Table of Results
    table_res = doc.add_table(rows=11, cols=5)
    table_res.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Metric / Dimension", "U-Net (Baseline)", "DeepLabV3+ (Proposed)", "Pix2Pix cGAN (Paarthu)", "CycleGAN (Shivam)"]
    for i, h in enumerate(headers):
        cell = table_res.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "2B6CB0")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    data_rows = [
        ("Student Allocation", "Shared Baseline", "Shared Proposed", "Paarthu (CB.SC.U4CSE23639)", "Shivam (CB.SC.U4CSE23445)"),
        ("Model Paradigm", "Supervised Discriminative", "Supervised Multi-Scale", "Supervised Paired cGAN", "Unpaired CycleGAN"),
        ("Generator / Model Params", "31,038,023 (~31.0M)", "3,935,783 (~3.9M)", "54,413,955 (~54.4M)", "7,841,411 (~7.8M)"),
        ("Total Architecture Params", "31,038,023", "3,935,783", "57,183,044", "21,214,856"),
        ("GPU Latency (RTX 4060)", "67.60 ms", "5.79 ms", "14.48 ms", "12.01 ms"),
        ("CPU Latency", "253.38 ms", "108.18 ms", "83.42 ms", "189.43 ms"),
        ("Throughput (GPU FPS)", "14.8 FPS", "172.7 FPS", "69.1 FPS", "83.3 FPS"),
        ("SSIM / PSNR Quality", "N/A (Mask Output)", "N/A (Mask Output)", "0.1719 / 12.10 dB", "0.1495 / 9.44 dB"),
        ("Laplacian Edge Sharpness", "N/A", "N/A", "4,443.1 (High Crispness)", "2,407.2 (Moderate)"),
        ("UNHCR Compliance Rate", "94.6% (Heuristic Solver)", "96.8% (Heuristic Solver)", "90.0% (Direct Synthesis)", "83.0% (Direct Synthesis)")
    ]

    for row_idx, r in enumerate(data_rows):
        row = table_res.rows[row_idx + 1]
        for col_idx, val in enumerate(r):
            cell = row.cells[col_idx]
            cell.text = val
            if row_idx % 2 == 1:
                set_cell_background(cell, "F7FAFC")
            if col_idx == 0:
                cell.paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()

    # 5. Visual Proof
    h1 = doc.add_heading("5. Visual Proof & Evaluation Figures", level=1)
    h1.style.font.color.rgb = RGBColor(26, 54, 93)

    fig1_path = os.path.join("backend", "static", "gan_evaluation_comparison.png")
    fig2_path = os.path.join("backend", "static", "gan_loss_curves.png")
    fig3_path = os.path.join("backend", "static", "gan_generated_samples.png")

    if os.path.exists(fig1_path):
        doc.add_heading("Figure 1: Hardware Latency, Parameters, and UNHCR Quality Comparison", level=2)
        doc.add_picture(fig1_path, width=Inches(6.2))
        p_cap = doc.add_paragraph("Fig 1: Quantitative benchmarking across Generator parameters, GPU latency, SSIM/PSNR, and UNHCR compliance.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if os.path.exists(fig2_path):
        doc.add_heading("Figure 2: Adversarial Training Loss Dynamics (Pix2Pix vs CycleGAN)", level=2)
        doc.add_picture(fig2_path, width=Inches(6.2))
        p_cap = doc.add_paragraph("Fig 2: Convergence trajectories of Pix2Pix (Paarthu) and CycleGAN (Shivam) during adaptation on the LoveDA dataset.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if os.path.exists(fig3_path):
        doc.add_heading("Figure 3: Qualitative Satellite-to-Blueprint Syntheses Grid", level=2)
        doc.add_picture(fig3_path, width=Inches(6.2))
        p_cap = doc.add_paragraph("Fig 3: Side-by-side visual matrix showing input satellite imagery, ground-truth benchmarks, Pix2Pix outputs, and CycleGAN outputs.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 6. Critical Discussion
    h1 = doc.add_heading("6. Critical Discussion & System Architecture Recommendations", level=1)
    h1.style.font.color.rgb = RGBColor(26, 54, 93)
    doc.add_paragraph(
        "1. Pix2Pix vs CycleGAN in Relief Operations: Pix2Pix produces sharply defined road lines and tent clusters (Laplacian score 4443.1 vs 2407.2) "
        "and achieves higher UNHCR compliance (90.0% vs 83.0%) when paired imagery is available. CycleGAN excels when deploying to new crisis zones "
        "without paired annotations, demonstrating superior noise robustness (45.0% vs 65.2% degradation).\n"
        "2. The Need for a Two-Stage Hybrid Architecture: Direct end-to-end GAN generation incurs a 2.7% to 7.0% hazard violation rate (placing tents "
        "on water bodies). In life-critical disaster relief, this is hazardous. The ideal architecture for ReliefPlan AI is a Two-Stage Pipeline: "
        "DeepLabV3+ performs deterministic hazard segmentation (5.79 ms, 0% water violations), followed by Pix2Pix/CycleGAN generating structured camp "
        "blueprints restricted strictly to safe buildable polygons."
    )

    # 7. Student Contribution Statement
    h1 = doc.add_heading("7. Individual Student Contributions", level=1)
    h1.style.font.color.rgb = RGBColor(26, 54, 93)
    doc.add_paragraph(
        "• Paarthu Reddy (CB.SC.U4CSE23639 - Team Lead): Engineered the Pix2Pix conditional GAN pipeline (gan_pix2pix.py), implementing the "
        "16-stage U-Net 256 generator with skip connections, 70x70 PatchGAN discriminator, L1 reconstruction loss, and SSIM/PSNR evaluation battery.\n"
        "• Shivam Kharat (CB.SC.U4CSE23445 - Core Researcher): Engineered the CycleGAN unpaired domain translation pipeline (gan_cyclegan.py), "
        "implementing dual ResNet-6 residual generators, dual unconditioned PatchGANs, cycle-consistency/identity loss formulation, and noise robustness benchmarking."
    )

    out_docx_path = os.path.join("documents", "lab_evaluation_2_report.docx")
    doc.save(out_docx_path)
    print(f"[*] Formatted report generated successfully at: {out_docx_path}")

if __name__ == "__main__":
    create_report_docx()
